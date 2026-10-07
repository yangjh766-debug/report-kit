# -*- coding: utf-8 -*-
"""기존 Word 자료의 서식을 재서 요약한다. 회사 양식에 맞춰 문서를 만들 때 먼저 돌린다.

사용: python3 docx_spec.py 기존자료.docx
출력: 쪽 여백, 머리글과 바닥글, 제목과 절 제목의 크기와 선, 본문 정렬과 줄간격, 대시 들여쓰기,
      캡션 형식, 표 괘선과 음영, 열 정렬, 셀 글자 크기. 값만 읽고 파일은 고치지 않는다.
"""
import re, sys, collections
import docx

tw = lambda v: round(v / 635) if v is not None else None   # EMU → twips


def borders(p):
    return re.findall(r'<w:(top|bottom) w:val="(\w+)" w:sz="(\d+)"[^>]*?w:color="(\w+)"', p._p.xml)


def main(path):
    d = docx.Document(path)
    s = d.sections[0]
    print(f"[쪽] {s.page_width.cm:.1f}x{s.page_height.cm:.1f}cm, 여백(twips) 상{tw(s.top_margin)} 하{tw(s.bottom_margin)} 좌{tw(s.left_margin)} 우{tw(s.right_margin)}")
    print(f"[머리글] {[p.text for p in s.header.paragraphs if p.text]}  [바닥글] {[p.text for p in s.footer.paragraphs if p.text]}")

    kinds = collections.OrderedDict()
    for p in d.paragraphs:
        t = p.text.strip()
        if not t and not borders(p):
            continue
        sizes = {r.font.size.pt for r in p.runs if r.font.size}
        bold = any(r.bold for r in p.runs)
        if re.match(r"^[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+\.", t): k = "절 제목(Ⅰ.)"
        elif re.match(r"^\[(표|그림) ?\d+\]", t): k = "캡션([표 N])"
        elif re.match(r"^\[.+\]$", t): k = "소제목([...])"
        elif re.match(r"^주\d*\)", t): k = "각주(주))"
        elif t.startswith("- "): k = "대시 항목"
        elif not t: k = "선만 있는 빈 줄"
        elif bold and sizes and max(sizes) >= 13: k = "문서 제목"
        else: k = "본문"
        if k in kinds: kinds[k]["n"] += 1; continue
        pf = p.paragraph_format
        kinds[k] = {"n": 1, "예": t[:30], "정렬": str(p.alignment).split(" ")[0] if p.alignment is not None else "상속",
                    "크기": sorted(sizes) or "기본", "굵게": bold, "앞뒤간격": (pf.space_before and pf.space_before.pt, pf.space_after and pf.space_after.pt),
                    "줄간격": pf.line_spacing, "들여쓰기": (tw(pf.left_indent), tw(pf.first_line_indent)), "선": borders(p)}
    print("\n[문단 종류]")
    for k, v in kinds.items():
        print(f"- {k} ({v.pop('n')}개): " + ", ".join(f"{a}={b}" for a, b in v.items()))

    print(f"\n[표] {len(d.tables)}개")
    for i, t in enumerate(d.tables[:5], 1):
        x = t._tbl.xml
        grid = [int(g) for g in re.findall(r'<w:gridCol w:w="(\d+)"', x)]
        edges = sorted(set(re.findall(r'<w:(top|bottom) w:val="(\w+)" w:sz="(\d+)"', x)))
        fills = sorted(set(re.findall(r'w:fill="(\w+)"', x)) - {"auto"})
        head = t.rows[0].cells
        al = []
        for j in range(len(t.columns)):
            vals = [r.cells[j].paragraphs[0].alignment for r in t.rows[1:] if j < len(r.cells)]
            al.append(collections.Counter(str(v).split(" ")[0] if v is not None else "왼쪽(상속)" for v in vals).most_common(1)[0][0] if vals else "-")
        sz = sorted({r.font.size.pt for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size})
        print(f"- 표{i}: {len(t.rows)}행 {len(t.columns)}열, 열폭 합 {sum(grid)}twips, 괘선 {edges}, 음영 {fills}, "
              f"머리행 반복 {'w:tblHeader' in x}, 글자 {sz}pt, 머리글 {[c.text[:8] for c in head]}, 열 정렬 {al}")


if __name__ == "__main__":
    main(sys.argv[1])
