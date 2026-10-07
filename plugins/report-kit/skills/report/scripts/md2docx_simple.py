import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.5)
st = doc.styles["Normal"]
st.font.name = "Malgun Gothic"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.3

HL = re.compile(r"(\[[^\]]*\])")
def add_runs(p, text, size=None):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        bold = part.startswith("**") and part.endswith("**")
        if bold: part = part[2:-2]
        for piece in HL.split(part):
            if not piece: continue
            r = p.add_run(piece)
            if bold: r.bold = True
            if size: r.font.size = size
            if HL.fullmatch(piece):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW

lines = open(src, encoding="utf-8").read().splitlines()
i = 0
first_title = True
while i < len(lines):
    ln = lines[i]
    if not ln.strip():
        i += 1; continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                rows.append(cells)
            i += 1
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c in range(ncol):
                cell = t.cell(r, c)
                cell.text = ""
                cp = cell.paragraphs[0]
                add_runs(cp, (row[c] if c < len(row) else ""), size=Pt(9.5))
                if r == 0:
                    for run in cp.runs: run.bold = True
        doc.add_paragraph()
        continue
    m = re.match(r"^(#+)\s+(.*)", ln)
    if m:
        level, text = len(m.group(1)), m.group(2)
        if level == 1 and first_title:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text); r.bold = True; r.font.size = Pt(16)
            first_title = False
        elif level == 1:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text); r.bold = True; r.font.size = Pt(13)
        else:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
            r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5)
        i += 1; continue
    m = re.match(r"^(\s*)(\d+\.|[가-힣]\.)\s+(.*)", ln)
    if m:
        indent = 0.6 if m.group(1) == "" else 1.4
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        add_runs(p, f"{m.group(2)} {m.group(3)}")
        i += 1; continue
    p = doc.add_paragraph()
    add_runs(p, ln.strip())
    i += 1
doc.save(dst)
print("saved", dst)
