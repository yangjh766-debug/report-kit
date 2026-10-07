# -*- coding: utf-8 -*-
"""pptx 글 상자마다 실제 글꼴(Pretendard)로 줄바꿈을 계산해 문제를 찾는다.

사용: python3 check_text.py 덱.pptx
찾는 것
- 단어 중간 끊김: 어절 단위 줄바꿈(eaLnBrk=0)이 꺼져 있는데 줄이 넘어가는 문단
- 외톨이 줄: 마지막 줄에 어절 하나만 남거나 줄 폭의 25%도 안 되는 경우
- 넘침: 계산한 줄 수 × 줄 높이가 글 상자 높이보다 큰 경우
Keynote와 PowerPoint의 실제 줄바꿈과 약간 다를 수 있으므로, 걸린 곳은 문구를 고치거나 의미 단위로 직접 줄을 나눈다.
"""
import re, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

FONT_DIRS = [Path.home() / "Library/Fonts", Path("/Library/Fonts"), Path("/usr/share/fonts")]


def font_file(name):
    for d in FONT_DIRS:
        for f in d.rglob(name) if d.exists() else []:
            return f
    sys.exit(f"글꼴 {name}을(를) 찾지 못했습니다. Pretendard를 설치하세요: https://github.com/orioncactus/pretendard")
_cache = {}


def font(face, bold, size):
    face = face or "Pretendard"
    w = "ExtraBold" if "ExtraBold" in face else "Black" if "Black" in face else "Bold" if bold else "Regular"
    key = (w, size)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(str(font_file(f"Pretendard-{w}.ttf")), int(size * 10))   # 0.1pt 단위로 재서 정밀도 확보
    return _cache[key]


def width_pt(text, face, bold, size):
    return font(face, bold, size).getlength(text) / 10


def tokens(p):
    """문단을 (어절, 글꼴, 굵게, 크기) 조각으로 나눈다. 글자 크기가 섞인 줄도 조각마다 따로 잰다."""
    toks = []
    for r in p.runs:
        size = r.font.size.pt if r.font.size else 18
        for i, part in enumerate(re.split(r"( )", r.text)):
            if part:
                toks.append((part, r.font.name, bool(r.font.bold), size))
    return toks


def wrap(toks, avail):
    lines, cur, cur_w = [], [], 0.0
    words, w = [], []
    for t in toks:                       # 공백 기준으로 어절 묶기(한 어절 안에 여러 run 가능)
        if t[0] == " ":
            if w: words.append(w); w = []
            words.append([t])
        else:
            w.append(t)
    if w: words.append(w)
    for word in words:
        ww = sum(width_pt(*t) for t in word)
        if word[0][0] == " ":
            if cur: cur.append(word); cur_w += ww
            continue
        if cur and cur_w + ww > avail:
            while cur and cur[-1][0][0] == " ": cur_w -= sum(width_pt(*t) for t in cur.pop())
            lines.append(cur); cur, cur_w = [], 0.0
        cur.append(word); cur_w += ww
    if cur: lines.append(cur)
    return ["".join(t[0] for wd in ln for t in wd) for ln in lines], [sum(width_pt(*t) for wd in ln for t in wd) for ln in lines]


def main(path):
    prs = Presentation(path); issues = 0
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.width is None:
                continue
            tf = sh.text_frame
            avail = Emu(sh.width).pt - Emu(tf.margin_left or 0).pt - Emu(tf.margin_right or 0).pt
            total_h = 0.0
            for p in tf.paragraphs:
                text = "".join(r.text for r in p.runs)
                if not text.strip():
                    continue
                r0 = p.runs[0]
                ea = p._p.pPr.get("eaLnBrk") if p._p.pPr is not None else None
                lines, widths = wrap(tokens(p), avail)
                size = max((r.font.size.pt if r.font.size else 18) for r in p.runs)
                spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                total_h += len(lines) * size * 1.2 * spacing
                if len(lines) > 1:
                    where = f"{si}장 [{text[:24]}…]"
                    if ea != "0":
                        print(f"{where} 단어 중간 끊김 가능: 어절 단위 줄바꿈이 꺼져 있음"); issues += 1
                    last = lines[-1]
                    if len(last.split(" ")) == 1 or widths[-1] < avail * 0.25:
                        print(f"{where} 외톨이 줄: 마지막 줄 「{last}」"); issues += 1
            box_h = Emu(sh.height).pt
            if total_h > box_h * 1.08 and total_h - box_h > 3:
                print(f"{si}장 [{tf.text[:24]}…] 넘침: 필요 {total_h:.0f}pt, 상자 {box_h:.0f}pt"); issues += 1
    print(f"검사 끝: 문제 {issues}건")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
